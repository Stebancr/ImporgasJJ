import base64
import json
import hashlib
import os
import shutil
import tempfile
from datetime import date, time, timedelta
from decimal import Decimal
from io import BytesIO
from urllib.parse import urlparse, parse_qs
from unittest.mock import patch

from django.core import mail
from django.core.management import call_command
from django.core.files.uploadedfile import SimpleUploadedFile
from django.core import signing
from django.test import override_settings
from django.urls import reverse
from django.utils import timezone
from pypdf import PdfReader
from PIL import Image as PILImage
from rest_framework.test import APITestCase

from usuarios.models import Credenciales, Usuario
from .models import ClienteVisita, VisitaTecnica, TipoVisita, CambioCostoVisita, ReporteVisita, EvidenciaFotografica
from .views import _save_pdf_and_notify, _persist_visit_pdf, _public_pdf_url
from crmChat.models import CRMContact, ChannelIntegration, ChatSession, ChatMessage
from .notifications import notify_technician_visit_assigned
from ecommerce.models import Notification


@override_settings(EMAIL_BACKEND='django.core.mail.backends.locmem.EmailBackend')
class VisitCompletionTests(APITestCase):
    def setUp(self):
        self.media_dir = tempfile.mkdtemp(prefix='visits-pdf-test-')
        self.media_settings = override_settings(MEDIA_ROOT=self.media_dir)
        self.media_settings.enable()
        self.addCleanup(shutil.rmtree, self.media_dir)
        self.addCleanup(self.media_settings.disable)
        profile = Usuario.objects.create(cedula='visit-admin', nombre_completo='Admin Visitas', correo='admin@example.com')
        self.admin = Credenciales.objects.create(usuario='visit-admin', usuario_rel=profile, estado=1, tipo_usuario=4)
        self.admin.set_password('safe-test-password')
        self.admin.save(update_fields=['password'])
        self.technician = Credenciales.objects.create(
            usuario='visit-technician', estado=1, tipo_usuario=1,
        )
        self.client.force_authenticate(self.admin)
        customer = ClienteVisita.objects.create(
            nombre='Cliente PDF', identificacion='123', telefono='3000000000',
            correo='cliente@example.com', direccion='Calle de prueba',
        )
        self.visit = VisitaTecnica.objects.create(
            cliente=customer, tipo_tarea='mantenimiento', fecha=date.today(),
            hora=time(10, 0), descripcion='Prueba de visita', creado_por=self.admin,
        )

    def valid_create_payload(self):
        return {
            'cliente_nombre': 'Cliente con valor',
            'cliente_identificacion': '1234567890',
            'cliente_telefono': '3001234567',
            'cliente_correo': 'cliente@example.com',
            'cliente_direccion': 'Calle 10 # 20-30',
            'tipo_tarea': 'revision_periodica',
            'fecha': date.today().isoformat(),
            'hora': '08:30:00',
            'valor_visita': 150000,
            'tecnico_id': self.technician.pk,
        }

    def test_task_number_is_deterministic(self):
        self.assertEqual(self.visit.numero_tarea, str(999 + self.visit.id))
        second = VisitaTecnica.objects.create(
            cliente=self.visit.cliente, tipo_tarea='revision', fecha=date.today(), hora=time(11, 0)
        )
        self.assertEqual(second.numero_tarea, str(999 + second.id))

    def test_search_document_phone_and_service_choice(self):
        for field, term in (('document', '123'), ('phone', '3000000000')):
            result = self.client.get(reverse('visitas-list-create'), {'search': term, 'search_field': field})
            self.assertEqual(result.status_code, 200)
            self.assertTrue(any(row['id'] == self.visit.pk for row in result.data))
        self.assertEqual(self.client.get(reverse('visitas-list-create'), {
            'search': 'abc', 'search_field': 'document',
        }).status_code, 400)
        payload = self.valid_create_payload()
        payload['tipo_tarea'] = 'not-an-option'
        self.assertEqual(self.client.post(reverse('visitas-list-create'), payload, format='json').status_code, 400)
        payload['tipo_tarea'] = 'revision_periodica'
        created = self.client.post(reverse('visitas-list-create'), payload, format='json')
        self.assertEqual(created.status_code, 201)
        self.assertEqual(created.data['tipo_tarea_display'], 'REVISION PERIODICA')
        detail = self.client.get(reverse('visitas-detail', args=[created.data['id']]))
        self.assertEqual(detail.data['tipo_tarea'], 'revision_periodica')

    def test_visit_value_is_created_and_returned_as_json_number(self):
        payload = self.valid_create_payload()
        created = self.client.post(reverse('visitas-list-create'), payload, format='json')
        self.assertEqual(created.status_code, 201)
        self.assertEqual(created.data['valor_visita'], 150000)
        self.assertEqual(created.data['costo_inicial'], 150000)
        self.assertEqual(created.data['costo_final'], 150000)
        self.assertIsInstance(created.data['valor_visita'], int)
        self.assertIsInstance(created.json()['valor_visita'], int)

        visit = VisitaTecnica.objects.get(pk=created.data['id'])
        self.assertEqual(visit.valor_visita, Decimal('150000.00'))
        self.assertEqual(visit.costo_inicial, Decimal('150000.00'))
        detail = self.client.get(reverse('visitas-detail', args=[visit.pk]))
        self.assertEqual(detail.status_code, 200)
        self.assertEqual(detail.data['valor_visita'], 150000)
        listed = self.client.get(reverse('visitas-list-create'))
        item = next(row for row in listed.data if row['id'] == visit.pk)
        self.assertEqual(item['valor_visita'], 150000)
        calendar = self.client.get(
            reverse('visitas-calendario'),
            {'mes': date.today().strftime('%Y-%m')},
        )
        calendar_item = next(
            row for rows in calendar.data.values() for row in rows if row['id'] == visit.pk
        )
        self.assertEqual(calendar_item['valor_visita'], 150000)

    def test_existing_visit_without_value_remains_compatible_and_negative_is_rejected(self):
        detail = self.client.get(reverse('visitas-detail', args=[self.visit.pk]))
        self.assertEqual(detail.status_code, 200)
        self.assertIsNone(detail.data['valor_visita'])

        invalid = self.client.patch(
            reverse('visitas-detail', args=[self.visit.pk]),
            {'valor_visita': -1},
            format='json',
        )
        self.assertEqual(invalid.status_code, 400)
        self.visit.refresh_from_db()
        self.assertIsNone(self.visit.valor_visita)

    def test_create_requires_operational_fields_but_allows_optional_id_and_cost(self):
        response = self.client.post(reverse('visitas-list-create'), {}, format='json')
        self.assertEqual(response.status_code, 400)
        required = {
            'cliente_nombre', 'cliente_telefono',
            'cliente_correo', 'cliente_direccion', 'tipo_tarea', 'fecha',
            'hora', 'tecnico_id',
        }
        self.assertTrue(required.issubset(response.data.keys()))
        self.assertNotIn('descripcion', response.data)
        self.assertNotIn('observaciones_iniciales', response.data)
        self.assertNotIn('cliente_identificacion', response.data)
        self.assertNotIn('valor_visita', response.data)
        self.assertEqual(str(response.data['tecnico_id'][0]), 'Debe seleccionar un técnico.')

        payload = self.valid_create_payload()
        payload.pop('cliente_identificacion')
        payload.pop('valor_visita')
        payload['cliente_indicaciones_llegada'] = 'Portería sur'
        created = self.client.post(reverse('visitas-list-create'), payload, format='json')
        self.assertEqual(created.status_code, 201)
        self.assertEqual(created.data['cliente']['identificacion'], '')
        self.assertEqual(created.data['cliente']['indicaciones_llegada'], 'Portería sur')
        self.assertIsNone(created.data['valor_visita'])
        self.assertEqual(created.data['creado_por'], self.admin.pk)

    def test_cost_update_by_assigned_technician_is_audited_and_exported(self):
        self.visit.tecnico = self.technician
        self.visit.save(update_fields=['tecnico'])
        self.client.force_authenticate(self.technician)
        response = self.client.patch(reverse('visitas-costo', args=[self.visit.pk]), {'valor_visita': 90000}, format='json')
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.data['valor_visita'], 90000)
        self.assertEqual(CambioCostoVisita.objects.get(visita=self.visit).usuario_id, self.technician.pk)
        self.assertEqual(self.client.patch(reverse('visitas-costo', args=[self.visit.pk]), {'valor_visita': -1}, format='json').status_code, 400)
        excel = self.client.get(reverse('visitas-exportar'), {'fecha_desde': date.today(), 'fecha_hasta': date.today()})
        self.assertEqual(excel.status_code, 200)
        from openpyxl import load_workbook
        rows = list(load_workbook(BytesIO(excel.content), read_only=True).active.values)
        self.assertIn(90000, rows[-1])

        other = ClienteVisita.objects.create(nombre='No asignado', direccion='Calle 3')
        forbidden_visit = VisitaTecnica.objects.create(
            cliente=other, tipo_tarea='revision_periodica',
            fecha=date.today(), hora=time(12),
        )
        self.assertEqual(self.client.patch(
            reverse('visitas-costo', args=[forbidden_visit.pk]),
            {'valor_visita': 100}, format='json',
        ).status_code, 404)
        self.assertEqual(self.client.post(
            reverse('visitas-tipos'), {'codigo': 'sin_permiso', 'nombre': 'Sin permiso'}, format='json',
        ).status_code, 403)

    def test_admin_partial_edit_requires_reason_and_preserves_initial_cost(self):
        created = self.client.post(reverse('visitas-list-create'), {
            **self.valid_create_payload(), 'valor_visita': 45000,
        }, format='json')
        self.assertEqual(created.status_code, 201)
        visit_id = created.data['id']
        url = reverse('visitas-detail', args=[visit_id])
        self.assertEqual(self.client.patch(url, {'valor_visita': 70000}, format='json').status_code, 400)
        visit = VisitaTecnica.objects.get(pk=visit_id)
        self.assertEqual(visit.valor_visita, Decimal('45000.00'))
        changed = self.client.patch(url, {
            'valor_visita': 70000, 'motivo_cambio_costo': 'Trabajo adicional',
        }, format='json')
        self.assertEqual(changed.status_code, 200, changed.data)
        visit.refresh_from_db()
        self.assertEqual(visit.costo_inicial, Decimal('45000.00'))
        self.assertEqual(visit.valor_visita, Decimal('70000.00'))
        self.assertEqual(CambioCostoVisita.objects.filter(visita=visit).count(), 2)

    def test_created_initial_cost_becomes_final_cost_after_technician_sync(self):
        payload = self.valid_create_payload()
        payload['valor_visita'] = 45000
        created = self.client.post(reverse('visitas-list-create'), payload, format='json')
        self.assertEqual(created.status_code, 201, created.data)
        visit_id = created.data['id']
        self.assertEqual(created.data['costo_inicial'], 45000)
        self.client.force_authenticate(self.technician)
        report = {
            'persona_atiende': 'Cliente de prueba', 'equipo': 'estufa',
            'ubicacion_equipo': 'cocina', 'motivo_servicio': 'Revisión',
            'solucion_realizada': 'Regulador cambiado', 'valor_servicio': '70000',
            'motivo_cambio_costo': 'Se agregó cambio de regulador',
            'fotos': [SimpleUploadedFile('evidencia.png', base64.b64decode('iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNk+A8AAQUBAScY42YAAAAASUVORK5CYII='), content_type='image/png')],
        }

        with self.captureOnCommitCallbacks(execute=True):
            finished = self.client.post(
                reverse('visitas-sincronizar', args=[visit_id]), report,
                format='multipart', HTTP_IDEMPOTENCY_KEY='final-cost-operation-001',
                HTTP_IF_MATCH=created.data['sync_version'],
            )
        self.assertEqual(finished.status_code, 200, finished.data)
        self.assertEqual(finished.data['costo_inicial'], 45000)
        self.assertEqual(finished.data['costo_final'], 70000)
        visit = VisitaTecnica.objects.get(pk=visit_id)
        self.assertEqual(visit.costo_inicial, Decimal('45000.00'))
        self.assertEqual(visit.valor_visita, Decimal('70000.00'))
        self.assertEqual(visit.reporte.valor_servicio, Decimal('70000.00'))
        self.assertIn('70,000', ''.join(page.extract_text() or '' for page in PdfReader(BytesIO(visit.pdf_final.read())).pages))
        self.client.force_authenticate(self.admin)
        listed = self.client.get(reverse('visitas-list-create'))
        self.assertEqual(next(row for row in listed.data if row['id'] == visit_id)['costo_final'], 70000)
        from openpyxl import load_workbook
        excel = self.client.get(reverse('visitas-exportar'), {'fecha_desde': date.today(), 'fecha_hasta': date.today()})
        self.assertIn(70000, [cell for row in load_workbook(BytesIO(excel.content), read_only=True).active.values for cell in row])
        cost_url = reverse('visitas-costo', args=[visit_id])
        amended = self.client.patch(cost_url, {
            'valor_visita': 80000, 'motivo_cambio_costo': 'Ajuste administrativo posterior',
        }, format='json')
        self.assertEqual(amended.status_code, 409, amended.data)
        visit.refresh_from_db()
        self.assertEqual(visit.costo_inicial, Decimal('45000.00'))
        self.assertEqual(visit.valor_visita, Decimal('70000.00'))
        self.assertTrue(self.client.get(reverse('visitas-detail', args=[visit_id])).data['pdf_disponible'])

    def test_assigned_technician_can_update_directions_without_erasing_other_fields(self):
        self.visit.tecnico = self.technician
        self.visit.save(update_fields=['tecnico'])
        url = reverse('visitas-indicaciones', args=[self.visit.pk])
        self.client.force_authenticate(self.technician)
        response = self.client.patch(url, {'indicaciones_llegada': 'Portería, torre B'}, format='json')
        self.assertEqual(response.status_code, 200, response.data)
        self.visit.cliente.refresh_from_db()
        self.assertEqual(self.visit.cliente.indicaciones_llegada, 'Portería, torre B')
        self.assertEqual(self.visit.cliente.direccion, 'Calle de prueba')

    def test_historical_evidence_is_retained_when_pdf_is_generated(self):
        photo = SimpleUploadedFile(
            'historica.png', base64.b64decode('iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNk+A8AAQUBAScY42YAAAAASUVORK5CYII='),
            content_type='image/png',
        )
        evidence = EvidenciaFotografica.objects.create(visita=self.visit, imagen=photo)
        ReporteVisita.objects.create(
            visita=self.visit, persona_atiende='Cliente', equipo='estufa',
            ubicacion_equipo='cocina', motivo_servicio='Revisión', solucion_realizada='Ajuste',
        )
        self.visit.estado = VisitaTecnica.ESTADO_FINALIZADA
        self.visit.save(update_fields=['estado'])
        detail = self.client.get(reverse('visitas-detail', args=[self.visit.pk]))
        self.assertTrue(detail.data['pdf_disponible'])
        with self.captureOnCommitCallbacks(execute=True):
            response = self.client.get(reverse('visitas-pdf', args=[self.visit.pk]))
        self.assertEqual(response.status_code, 200)
        response.close()
        self.visit.refresh_from_db()
        self.assertTrue(self.visit.pdf_final)
        self.assertEqual(len(mail.outbox), 0)
        evidence.refresh_from_db()
        self.assertFalse(evidence.es_temporal)
        self.assertTrue(evidence.imagen.storage.exists(evidence.imagen.name))
        self.assertIsNone(evidence.eliminada_en)

    def test_cleanup_only_removes_old_unreferenced_temporary_files(self):
        folder = os.path.join(self.media_dir, 'evidencias_temporales')
        os.makedirs(folder, exist_ok=True)
        orphan = os.path.join(folder, 'orphan.png')
        with open(orphan, 'wb') as output:
            output.write(b'synthetic test file')
        old = timezone.now().timestamp() - 32 * 86400
        os.utime(orphan, (old, old))
        photo = SimpleUploadedFile('retained.png', base64.b64decode(
            'iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNk+A8AAQUBAScY42YAAAAASUVORK5CYII='
        ), content_type='image/png')
        evidence = EvidenciaFotografica.objects.create(visita=self.visit, imagen=photo, es_temporal=True)
        retained = evidence.imagen.path
        os.utime(retained, (old, old))
        call_command('cleanup_visit_temporary_images', days=30)
        self.assertTrue(os.path.exists(orphan))
        call_command('cleanup_visit_temporary_images', days=30, apply=True)
        self.assertFalse(os.path.exists(orphan))
        self.assertTrue(os.path.exists(retained))

    def test_excel_filters_and_escapes_user_text(self):
        self.visit.cliente.nombre = '=SUM(1,1)'
        self.visit.cliente.save(update_fields=['nombre'])
        other = ClienteVisita.objects.create(nombre='Otro cliente', direccion='Calle 2', telefono='3000000001')
        VisitaTecnica.objects.create(
            cliente=other, tecnico=self.technician, tipo_tarea='revision_periodica',
            fecha=date.today(), hora=time(11), creado_por=self.admin,
        )
        params = {'fecha_desde': date.today(), 'fecha_hasta': date.today()}
        from openpyxl import load_workbook
        response = self.client.get(reverse('visitas-exportar'), params)
        rows = list(load_workbook(BytesIO(response.content), read_only=True).active.values)
        self.assertEqual(len(rows), 3)
        self.assertEqual(rows[1][4], "'=SUM(1,1)")
        self.assertEqual(rows[1][10], 'No disponible')
        self.assertEqual(self.client.get(reverse('visitas-exportar'), {**params, 'estado': 'invalido'}).status_code, 400)
        filtered = self.client.get(reverse('visitas-exportar'), {
            **params, 'tecnico_id': self.technician.pk, 'tipo_tarea': 'revision_periodica',
        })
        rows = list(load_workbook(BytesIO(filtered.content), read_only=True).active.values)
        self.assertEqual(len(rows), 2)
        self.assertEqual(rows[1][4], 'Otro cliente')

    def test_types_preserve_old_labels_and_can_be_deactivated(self):
        items = self.client.get(reverse('visitas-tipos'))
        self.assertEqual(items.status_code, 200)
        self.assertTrue(any(item['codigo'] == 'revision_periodica' for item in items.data))
        created = self.client.post(reverse('visitas-tipos'), {'codigo': 'nueva_visita', 'nombre': 'Nueva visita'}, format='json')
        self.assertEqual(created.status_code, 201)
        item = TipoVisita.objects.get(codigo='nueva_visita')
        self.assertEqual(self.client.patch(reverse('visitas-tipo-detail', args=[item.pk]), {'activo': False}, format='json').status_code, 200)
        self.assertEqual(self.client.post(reverse('visitas-list-create'), {**self.valid_create_payload(), 'tipo_tarea': item.codigo}, format='json').status_code, 400)

    @override_settings(FRONTEND_PUBLIC_URL='http://localhost')
    def test_admin_can_complete_visit_with_photo_without_signature_and_pdf_is_protected(self):
        payload = self.valid_create_payload()
        payload.pop('cliente_identificacion')
        payload.pop('valor_visita')
        payload['completar_ahora'] = 'true'
        payload['reporte'] = json.dumps({
            'persona_atiende': 'Cliente de prueba', 'equipo': 'estufa',
            'ubicacion_equipo': 'cocina', 'motivo_servicio': 'Revisión',
            'solucion_realizada': 'Equipo revisado',
        })
        payload['fotos'] = SimpleUploadedFile(
            'foto.png', base64.b64decode('iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNk+A8AAQUBAScY42YAAAAASUVORK5CYII='),
            content_type='image/png',
        )
        with self.captureOnCommitCallbacks(execute=True):
            created = self.client.post(reverse('visitas-list-create'), payload, format='multipart')
        self.assertEqual(created.status_code, 201, created.data)
        self.assertNotIn('pdf_final', created.data)
        visit = VisitaTecnica.objects.get(pk=created.data['id'])
        self.assertEqual(visit.estado, VisitaTecnica.ESTADO_FINALIZADA)
        self.assertEqual(visit.evidencias.count(), 1)
        self.assertTrue(visit.pdf_final)
        self.assertEqual(visit.whatsapp_notificacion_estado, 'manual')
        with override_settings(DEBUG=True):
            self.assertTrue(_public_pdf_url(visit).startswith('http://localhost/api/visits/'))
        with override_settings(DEBUG=False):
            self.assertIsNone(_public_pdf_url(visit))
        reader = PdfReader(BytesIO(visit.pdf_final.read()))
        self.assertGreater(len(reader.pages[0].images), 0, 'El logo no se incorporó al PDF.')
        text = ''.join(page.extract_text() or '' for page in reader.pages)
        self.assertIn('Visita finalizada sin firma del cliente', text)
        self.assertIn('Admin Visitas', text)
        self.assertNotIn('Indicaciones para llegar', text)
        self.assertEqual(self.client.get('/media/' + visit.pdf_final.name).status_code, 404)
        self.assertEqual(self.client.get(reverse('visitas-pdf-publico', args=[visit.pk])).status_code, 403)

    def test_admin_completion_without_images_still_stores_valid_pdf(self):
        payload = self.valid_create_payload()
        payload['completar_ahora'] = 'true'
        payload['reporte'] = json.dumps({
            'persona_atiende': 'Cliente sintético', 'equipo': 'estufa',
            'ubicacion_equipo': 'cocina', 'motivo_servicio': 'Revisión',
            'solucion_realizada': 'Trabajo terminado',
        })
        with self.captureOnCommitCallbacks(execute=True):
            created = self.client.post(reverse('visitas-list-create'), payload, format='multipart')
        self.assertEqual(created.status_code, 201, created.data)
        visit = VisitaTecnica.objects.get(pk=created.data['id'])
        self.assertEqual(visit.estado, VisitaTecnica.ESTADO_FINALIZADA)
        self.assertEqual(visit.evidencias.count(), 0)
        self.assertTrue(visit.pdf_final.storage.exists(visit.pdf_final.name))
        self.assertIn('Visita finalizada sin firma del cliente', ''.join(
            page.extract_text() or '' for page in PdfReader(BytesIO(visit.pdf_final.read())).pages))

    @override_settings(FRONTEND_PUBLIC_URL='https://www.imporgasjj.com')
    def test_definitive_pdf_has_safe_name_all_images_metadata_and_revocable_links(self):
        def photo(name, color):
            image = BytesIO()
            PILImage.new('RGB', (120, 80), color).save(image, format='PNG')
            return SimpleUploadedFile(name, image.getvalue(), content_type='image/png')

        payload = self.valid_create_payload()
        payload['cliente_nombre'] = 'Juan Pérez / prueba'
        payload['completar_ahora'] = 'true'
        payload['reporte'] = json.dumps({
            'persona_atiende': 'Juan', 'equipo': 'estufa', 'ubicacion_equipo': 'cocina',
            'motivo_servicio': 'Revisión', 'solucion_realizada': 'Trabajo terminado',
        })
        payload['fotos'] = [photo('roja.png', 'red'), photo('azul.png', 'blue')]
        with self.captureOnCommitCallbacks(execute=True):
            created = self.client.post(reverse('visitas-list-create'), payload, format='multipart')
        self.assertEqual(created.status_code, 201, created.data)
        visit = VisitaTecnica.objects.get(pk=created.data['id'])
        self.assertEqual(visit.pdf_estado, 'generado')
        self.assertRegex(visit.pdf_final.name, r'/VISITA-\d{6}-Juan-Perez-prueba\.pdf$')
        pdf_bytes = visit.pdf_final.read()
        self.assertEqual(visit.pdf_bytes, len(pdf_bytes))
        self.assertEqual(visit.pdf_sha256, hashlib.sha256(pdf_bytes).hexdigest())
        self.assertIsNotNone(visit.pdf_generado_en)
        self.assertEqual(sum(len(page.images) for page in PdfReader(BytesIO(pdf_bytes)).pages), 3)
        for evidence in visit.evidencias.all():
            self.assertEqual(evidence.imagen.name, '')
            self.assertIsNotNone(evidence.eliminada_en)
        admin_pdf = self.client.get(reverse('visitas-pdf', args=[visit.pk]))
        self.assertEqual(admin_pdf.status_code, 200)
        self.assertIn('inline', admin_pdf['Content-Disposition'])
        self.assertIn('Juan-Perez-prueba.pdf', admin_pdf['Content-Disposition'])
        outsider = Credenciales.objects.create(usuario='qa-other-technician', tipo_usuario=1, estado=1)
        self.client.force_authenticate(outsider)
        self.assertEqual(self.client.get(reverse('visitas-pdf', args=[visit.pk])).status_code, 404)
        self.client.force_authenticate(self.admin)
        detail = self.client.get(reverse('visitas-detail', args=[visit.pk]))
        self.assertEqual(detail.data['evidencias'], [])
        self.assertFalse(detail.data['costo_editable'])
        pdf_token = signing.dumps({'visit_id': visit.pk, 'version': 1}, salt='visita-pdf-publico')
        self.assertEqual(self.client.get(reverse('visitas-pdf-publico', args=[visit.pk]), {'token': pdf_token}).status_code, 200)
        with patch('django.core.signing.time.time', return_value=1):
            expired_token = signing.dumps({'visit_id': visit.pk, 'version': 1}, salt='visita-pdf-publico')
        self.assertEqual(self.client.get(reverse('visitas-pdf-publico', args=[visit.pk]), {'token': expired_token}).status_code, 403)
        self.assertEqual(self.client.post(reverse('visitas-pdf-revocar', args=[visit.pk])).status_code, 200)
        self.assertEqual(self.client.get(reverse('visitas-pdf-publico', args=[visit.pk]), {'token': pdf_token}).status_code, 403)
        self.assertEqual(self.client.get(reverse('visitas-pdf-publico', args=[self.visit.pk]), {'token': pdf_token}).status_code, 403)
        stored_name = visit.pdf_final.name
        _save_pdf_and_notify(visit.pk)
        visit.refresh_from_db()
        self.assertEqual(visit.pdf_final.name, stored_name)
        self.assertEqual(self.client.post(reverse('visitas-fotos', args=[visit.pk]), {
            'fotos': photo('tardia.png', 'green'),
        }, format='multipart').status_code, 409)
        self.assertEqual(self.client.delete(reverse('visitas-fotos-delete', args=[visit.pk, visit.evidencias.first().pk])).status_code, 409)
        from openpyxl import load_workbook
        export = self.client.get(reverse('visitas-exportar'), {
            'fecha_desde': date.today(), 'fecha_hasta': date.today(), 'estado': 'finalizada',
        })
        workbook = load_workbook(BytesIO(export.content))
        sheet = workbook['Visitas']
        self.assertIn('VisitasTecnicas', sheet.tables)
        self.assertEqual(sheet.freeze_panes, 'A2')
        self.assertEqual(sheet['A1'].fill.fgColor.rgb[-6:], '173A57')
        self.assertEqual(sheet['B2'].number_format, 'dd/mm/yyyy')
        self.assertIn('#,##0.00', sheet['I2'].number_format)
        self.assertEqual(sheet['K2'].value, 'Ver PDF')
        self.assertTrue(sheet['K2'].hyperlink.target.startswith('https://www.imporgasjj.com/api/visits/'))
        excel_url = urlparse(sheet['K2'].hyperlink.target)
        self.assertEqual(self.client.get(excel_url.path.removeprefix('/api'), {
            'token': parse_qs(excel_url.query)['token'][0],
        }).status_code, 200)
        summary = workbook['Resumen']
        self.assertEqual(summary['B5'].value, 'Finalizada')
        self.assertEqual(summary['D5'].value, 1)

    def test_pdf_failure_keeps_original_images_and_can_be_retried(self):
        payload = self.valid_create_payload()
        payload['completar_ahora'] = 'true'
        payload['reporte'] = json.dumps({
            'persona_atiende': 'Cliente', 'equipo': 'estufa', 'ubicacion_equipo': 'cocina',
            'motivo_servicio': 'Revisión', 'solucion_realizada': 'Trabajo terminado',
        })
        payload['fotos'] = SimpleUploadedFile(
            'evidencia.png', base64.b64decode('iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNk+A8AAQUBAScY42YAAAAASUVORK5CYII='),
            content_type='image/png',
        )
        with patch('AppVisits.views._generar_pdf', side_effect=ValueError('fallo de prueba')):
            with self.captureOnCommitCallbacks(execute=True):
                created = self.client.post(reverse('visitas-list-create'), payload, format='multipart')
        self.assertEqual(created.status_code, 201)
        visit = VisitaTecnica.objects.get(pk=created.data['id'])
        evidence = visit.evidencias.get()
        self.assertEqual(visit.pdf_estado, 'fallido')
        self.assertNotEqual(visit.estado, VisitaTecnica.ESTADO_FINALIZADA)
        self.assertFalse(visit.pdf_final)
        self.assertIsNone(evidence.archivada_en)
        self.assertTrue(evidence.imagen.storage.exists(evidence.imagen.name))
        self.assertEqual(self.client.get(reverse('visitas-pdf', args=[visit.pk])).status_code, 409)
        with self.captureOnCommitCallbacks(execute=True):
            retried = self.client.post(reverse('visitas-pdf-reintentar', args=[visit.pk]))
        self.assertEqual(retried.status_code, 200, retried.data)
        visit.refresh_from_db()
        evidence.refresh_from_db()
        self.assertEqual(visit.pdf_estado, 'generado')
        self.assertEqual(visit.estado, VisitaTecnica.ESTADO_FINALIZADA)
        self.assertIsNotNone(evidence.eliminada_en)
        self.assertEqual(evidence.imagen.name, '')

    @override_settings(FRONTEND_PUBLIC_URL='https://www.imporgasjj.com')
    @patch('crmChat.apps.meta.services.dispatch_outbound_message')
    def test_pdf_whatsapp_only_uses_matching_recent_conversation(self, send):
        self.visit.estado = VisitaTecnica.ESTADO_FINALIZADA
        self.visit.cliente.telefono = '3001234567'
        self.visit.cliente.save(update_fields=['telefono'])
        self.visit.save(update_fields=['estado'])
        ReporteVisita.objects.create(
            visita=self.visit, persona_atiende='Cliente', equipo='estufa',
            ubicacion_equipo='cocina', motivo_servicio='Revisar', solucion_realizada='Revisado',
        )
        integration = ChannelIntegration.objects.create(
            name='Prueba', channel='whatsapp_web', active=True,
            connection_status='connected', external_account_id='test-only',
        )
        wrong = CRMContact.objects.create(name='Otro', phone='3007654321')
        session = ChatSession.objects.create(
            user_name='Otro', contact=wrong, integration=integration,
            channel='whatsapp_web', external_thread_id='573007654321@s.whatsapp.net',
            status='active', last_customer_message_at=timezone.now(),
        )
        ChatMessage.objects.create(session=session, text='Hola', sender_type='user', direction='inbound', external_message_id='inbound-test-1')
        _save_pdf_and_notify(self.visit.pk)
        self.visit.refresh_from_db()
        self.assertEqual(self.visit.whatsapp_notificacion_estado, 'manual')
        send.assert_not_called()

        correct = CRMContact.objects.create(name='Cliente', phone='3001234567')
        session.contact = correct
        session.external_thread_id = '573001234567@s.whatsapp.net'
        session.save(update_fields=['contact', 'external_thread_id'])
        with override_settings(FRONTEND_PUBLIC_URL='https://backend'):
            _save_pdf_and_notify(self.visit.pk)
        self.visit.refresh_from_db()
        self.assertEqual(self.visit.whatsapp_notificacion_estado, 'manual')
        send.assert_not_called()
        _save_pdf_and_notify(self.visit.pk)
        send.assert_called_once()
        self.visit.refresh_from_db()
        self.assertEqual(self.visit.whatsapp_notificacion_estado, 'enviada')
        message = ChatMessage.objects.get(client_message_id=f'visit-pdf-{self.visit.pk}')
        self.assertIn('https://www.imporgasjj.com/api/visits/', message.text)
        token = signing.dumps({'visit_id': self.visit.pk}, salt='visita-pdf-publico')
        response = self.client.get(reverse('visitas-pdf-publico', args=[self.visit.pk]), {'token': token})
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response['Content-Type'], 'application/pdf')
        with self.captureOnCommitCallbacks(execute=True):
            updated = self.client.patch(reverse('visitas-costo', args=[self.visit.pk]), {'valor_visita': 12345}, format='json')
        self.assertEqual(updated.status_code, 200)
        self.visit.refresh_from_db()
        revised_text = ''.join(page.extract_text() or '' for page in PdfReader(BytesIO(self.visit.pdf_final.read())).pages)
        self.assertIn('12,345', revised_text)
        send.assert_called_once()

    def test_create_rejects_invalid_identity_phone_email_past_date_and_technician(self):
        cases = [
            ('cliente_identificacion', '123ABC', 'La cédula solo puede contener números.'),
            ('cliente_telefono', '300-123-4567', 'El número de celular solo puede contener números.'),
            ('cliente_correo', 'usuario@dominio', 'Ingrese un correo electrónico válido.'),
            ('fecha', (date.today() - timedelta(days=1)).isoformat(), 'La fecha de la visita no puede ser anterior al día actual.'),
            ('tecnico_id', self.admin.pk, 'Debe seleccionar un técnico válido y activo.'),
        ]
        for field, value, expected in cases:
            with self.subTest(field=field):
                payload = self.valid_create_payload()
                payload[field] = value
                response = self.client.post(reverse('visitas-list-create'), payload, format='json')
                self.assertEqual(response.status_code, 400)
                self.assertEqual(str(response.data[field][0]), expected)

    def test_finalized_visit_cannot_be_patched_but_pending_visit_can(self):
        url = reverse('visitas-detail', args=[self.visit.pk])
        editable = self.client.patch(url, {'descripcion': 'Todavía editable'}, format='json')
        self.assertEqual(editable.status_code, 200)

        self.visit.refresh_from_db()
        self.visit.estado = VisitaTecnica.ESTADO_EN_PROCESO
        self.visit.save(update_fields=['estado', 'fecha_actualizacion'])
        in_progress = self.client.patch(url, {'descripcion': 'En proceso editable'}, format='json')
        self.assertEqual(in_progress.status_code, 200)

        self.visit.refresh_from_db()
        self.visit.estado = VisitaTecnica.ESTADO_FINALIZADA
        self.visit.save(update_fields=['estado', 'fecha_actualizacion'])
        blocked = self.client.patch(url, {'descripcion': 'Cambio no permitido'}, format='json')
        self.assertEqual(blocked.status_code, 400)
        self.assertEqual(blocked.data['detail'], 'No se puede modificar una visita que ya está finalizada.')
        self.visit.refresh_from_db()
        self.assertEqual(self.visit.descripcion, 'En proceso editable')

    def test_completion_sends_one_email_with_pdf_and_signature_line_label(self):
        signature_png = base64.b64encode(base64.b64decode(
            'iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNk+A8AAQUBAScY42YAAAAASUVORK5CYII='
        )).decode()
        payload = {
            'persona_atiende': 'Cliente PDF', 'equipo': 'estufa',
            'ubicacion_equipo': 'cocina', 'motivo_servicio': 'Mantenimiento',
            'solucion_realizada': 'Trabajo completado',
            'firma_base64': f'data:image/png;base64,{signature_png}',
        }
        url = reverse('visitas-finalizar', args=[self.visit.id])
        with self.captureOnCommitCallbacks(execute=True):
            response = self.client.post(url, payload, format='json')
        self.assertEqual(response.status_code, 200)
        self.visit.refresh_from_db()
        self.assertTrue(self.visit.pdf_final)
        self.visit.refresh_from_db()
        self.assertIsNotNone(self.visit.correo_completada_en)
        self.assertEqual(len(mail.outbox), 1)
        attachment = mail.outbox[0].attachments[0]
        self.assertEqual(attachment.mimetype, 'application/pdf')
        with self.visit.pdf_final.open('rb') as saved_pdf:
            self.assertEqual(attachment.content, saved_pdf.read())
        text = ''.join(page.extract_text() or '' for page in PdfReader(BytesIO(attachment.content)).pages)
        self.assertIn(f'Tarea {self.visit.numero_tarea}', text)
        self.assertIn('Firma cliente', text)

        repeated = self.client.post(url, payload, format='json')
        self.assertEqual(repeated.status_code, 400)
        self.assertEqual(len(mail.outbox), 1)


class VisitAssignmentNotificationTests(APITestCase):
    @patch('AppVisits.notifications.send_push_to_user')
    def test_assigned_technician_receives_persisted_and_fcm_notification(self, send_push):
        technician = Credenciales.objects.create(usuario='tech-push', estado=1, tipo_usuario=1)
        customer = ClienteVisita.objects.create(nombre='Cliente Push', direccion='Dirección')
        visit = VisitaTecnica.objects.create(
            cliente=customer, tecnico=technician, tipo_tarea='revision',
            fecha=date.today(), hora=time(9, 0), valor_visita=Decimal('150000.00'),
        )
        send_push.return_value = {'sent': 1, 'failed': 0, 'deactivated': 0}

        result = notify_technician_visit_assigned(visit)

        self.assertEqual(result['sent'], 1)
        notification = Notification.objects.get(user=technician)
        self.assertIn(visit.numero_tarea, notification.message)
        send_push.assert_called_once()
        self.assertEqual(send_push.call_args.args[3]['valor_visita'], '150000.00')
